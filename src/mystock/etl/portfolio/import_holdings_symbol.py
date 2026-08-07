"""Imports a broker holdings export keyed by plain NSE symbol (Zerodha-style: sheet with
columns 'Symbol (N)', 'Category', 'Net Qty', 'Avg. Price', 'LTP', 'Current Value', ...) as
OPENING_BALANCE transactions, using the broker's own reported Avg. Price as cost basis.
This broker's export has no trade history and no ISIN, only the symbol directly (which
in practice matches the NSE trading symbol used in main.dim_security).

Only 'NSE EQ' category rows are imported; anything else (e.g. a pledged/T1 sub-sheet
duplicate, or a non-equity category) is skipped and reported.

Usage:
    python -m mystock.etl.portfolio.import_holdings_symbol --file ~/Downloads/Holdings_04-Aug-2026.xlsx \\
        --user raj --portfolio "Zerodha" --broker zerodha --as-of 2026-08-04
"""
import argparse
from datetime import datetime
from pathlib import Path

import openpyxl

from mystock.db import get_conn
from mystock.etl.portfolio.db_helpers import get_or_create_user, get_or_create_portfolio, match_security_by_symbol


def to_float(v):
    if v is None:
        return 0.0
    if isinstance(v, str):
        v = v.replace(",", "").strip()
    return float(v or 0)


def find_header_row(rows):
    for i, r in enumerate(rows):
        if r and len(r) > 2 and r[1] == "Category" and r[2] == "Net Qty":
            return i
    raise ValueError("could not find header row (expected 'Category', 'Net Qty' columns)")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--file", required=True)
    ap.add_argument("--sheet", default=None, help="defaults to the first sheet in the workbook")
    ap.add_argument("--user", required=True)
    ap.add_argument("--portfolio", required=True)
    ap.add_argument("--broker", default=None)
    ap.add_argument("--as-of", required=True, help="report date (YYYY-MM-DD) — this file has no per-row date")
    args = ap.parse_args()

    txn_date = datetime.strptime(args.as_of, "%Y-%m-%d").date()

    wb = openpyxl.load_workbook(Path(args.file).expanduser(), data_only=True)
    ws = wb[args.sheet] if args.sheet else wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    header_idx = find_header_row(rows)
    data_rows = [r for r in rows[header_idx + 1:] if r and r[0]]

    conn = get_conn()
    user_id = get_or_create_user(conn, args.user)
    portfolio_id = get_or_create_portfolio(conn, user_id, args.portfolio, broker=args.broker)

    inserted = skipped = 0
    skipped_rows = []

    with conn.cursor() as cur:
        for r in data_rows:
            symbol, category, net_qty, avg_price = str(r[0]).strip(), r[1], r[2], r[3]
            if category != "NSE EQ":
                skipped += 1
                skipped_rows.append((symbol, category, "not an NSE EQ row"))
                continue
            qty = to_float(net_qty)
            if qty <= 0:
                continue
            match = match_security_by_symbol(conn, symbol)
            if not match:
                skipped += 1
                skipped_rows.append((symbol, category, "symbol not in main.dim_security"))
                continue
            security_key, _ = match
            cur.execute(
                """INSERT INTO portfolio.transactions
                   (portfolio_id, security_key, transaction_type, transaction_date, quantity, price, source_file)
                   VALUES (%s, %s, 'OPENING_BALANCE', %s, %s, %s, %s)
                   ON CONFLICT (portfolio_id, security_key, transaction_type, transaction_date, quantity, price)
                   DO NOTHING""",
                (portfolio_id, security_key, txn_date, qty, to_float(avg_price), Path(args.file).name),
            )
            inserted += 1

    conn.commit()
    conn.close()

    print(f"portfolio_id={portfolio_id}: {inserted} equity opening balances inserted, {skipped} skipped")
    for symbol, category, reason in skipped_rows:
        print(f"  SKIPPED {symbol} ({category}): {reason}")


if __name__ == "__main__":
    main()
